"""Fill question-per-row official forms from a Markdown answer file.

Many committee forms are one big table, one question per row, in a handful of
shapes. Instead of a hand-written generator per form, `scan()` reads the blank
and classifies every row, so a pack only declares *which* blanks it fills:

  header  label | (empty cell)          計畫主持人 | ＿     → value from config.toml
  yesno   question | □ | □ | reason      columns named by the row above (是 | 否)
  choice  question … □ option □ option   (also 'option□' with the box after it)
  text    question | (empty cell)        or one cell: heading + guidance paragraphs

Each row gets a key `<FORM>-<n>` (AF127-6). Answers come from the study's
answer file (`attachments = "@附件填寫.md"`), one `## <key> …` per row:

  yesno   是 / 否 / 是：理由              → ticks the column, reason into the last cell
  choice  - option prefix (or `- option：說明` to also fill its ＿＿)
  text    the text (lines → paragraphs); in a one-cell row it replaces the guidance

`skeleton()` lists every row of a blank as Markdown so a study can answer it;
unanswered rows are left exactly as the blank has them.
"""
import copy
import re
from dataclasses import dataclass, field

from docx.table import Table

from rec_helper.template_fill import UNCHECKED, Region, cell_paragraphs, check

BOX = "□☐"
BOX_RE = re.compile(f"[{BOX}]")
OPTION_AFTER_BOX = re.compile(rf"[{BOX}]\s*([^{BOX}\n]*?)\s*(?=[{BOX}\n]|$)")
OPTION_BEFORE_BOX = re.compile(rf"([^\s{BOX}]+)\s*[{BOX}]")
HEADER_LABELS = {          # first cell → config value (None: leave for the study to answer)
    "計畫主持人": "pi.name", "試驗主持人": "pi.name", "主要主持人": "pi.name", "主持人": "pi.name",
    "中文": "study.title_zh", "英文": "study.title_en",
    "計畫名稱": "study.title_zh", "中文計畫名稱": "study.title_zh", "計畫中文名稱": "study.title_zh",
    "計畫英文名稱": "study.title_en", "研究倫理委員會編號": "study.irb_no",
}
def _pi_title(config):
    """cv.toml dept '影像醫學部／醫事放射師' → '醫事放射師'."""
    return str(config.get("pi", {}).get("dept", "")).partition("／")[2].strip()


def _co_pis(config):
    return "、".join(p.get("name", "") for p in config.get("co_pi") or [] if isinstance(p, dict))


INLINE_LABELS = (          # 'label：' inside a header paragraph → config path or callable(config)
    ("中文：", "study.title_zh"), ("英文：", "study.title_en"),
    ("試驗主持人：", "pi.name"), ("主要主持人：", "pi.name"),
    ("職稱：", _pi_title),           # the first 職稱 of a header row is the PI's
    ("協同主持人：", _co_pis),
)


@dataclass
class Item:
    key: str
    kind: str                       # header | yesno | choice | text
    question: str
    cells: list                     # unique cells of the row
    options: tuple = ()
    columns: dict = field(default_factory=dict)   # yesno: label → cell index
    answer_cell: object = None      # text/header: the empty cell to write in
    source: str = ""                # header: dotted config path
    paras: list = None              # choice outside any table: the body paragraphs

    def paragraphs(self):
        return self.paras if self.paras is not None else [p for c in self.cells for p in cell_paragraphs(c)]


def _unique_cells(row):
    seen, out = set(), []
    for c in row.cells:
        if c._tc not in seen:
            seen.add(c._tc)
            out.append(c)
    return out


def _clean(text):
    return re.sub(r"\s+", " ", text).strip()


def _is_box(text):
    """A cell holding only checkboxes (one per sub-question line)."""
    t = re.sub(r"\s", "", text)
    return bool(t) and all(ch in BOX for ch in t)


def _box_after(text):
    """'是□ 否□ N/A□' (label, then its box): every line with boxes ends in a box."""
    lines = [ln for ln in text.split("\n") if BOX_RE.search(ln)]
    return bool(lines) and all(re.search(rf"[{BOX}]\s*$", ln) and not re.match(rf"\s*[{BOX}]", ln)
                               for ln in lines)


def _options(text):
    if _box_after(text):
        return tuple(o.strip(" _＿：:")[:30] for o in OPTION_BEFORE_BOX.findall(text))   # 是□ 否□
    return tuple(o.strip(" _＿：:")[:60] for o in OPTION_AFTER_BOX.findall(text) if o.strip(" _＿：:"))


def _top_tables(doc):
    """The body's top-level tables, in order (nested tables stay inside their cell)."""
    for block in doc.element.body.iterchildren():
        if block.tag.rsplit("}", 1)[-1] == "tbl":
            yield Table(block, doc)


def _is_list(table):
    """A log table (序號 | 日期 | 事件 …): a header row, then rows left empty to be filled."""
    rows = [[_clean(c.text) for c in _unique_cells(r)] for r in table.rows]
    if len(rows) < 3 or len(rows[0]) < 3:
        return False
    empty = sum(1 for r in rows[1:] if len(r) == len(rows[0]) and not any(r))
    return empty >= 2 or (empty >= 1 and len(rows) <= 4)


def _list_columns(table):
    """Column titles from the header row(s): first line of each cell."""
    return tuple(_clean(c.text.split("\n")[0]) for c in _unique_cells(table.rows[0]))


def _rows(doc, form_id, items, counter):
    """Rows of non-list tables; list tables become one `list` item each."""
    for table in _top_tables(doc):
        if _is_list(table):
            counter[0] += 1
            items.append(Item(f"{form_id}-{counter[0]}", "list", "／".join(_list_columns(table)), [],
                              _list_columns(table), paras=table))
            continue
        yield from table.rows


def scan(doc, form_id):
    items, header_labels, counter = [], None, [0]
    for row in _rows(doc, form_id, items, counter):
        n = counter[0]
        cells = _unique_cells(row)
        texts = [_clean(c.text) for c in cells]
        if not any(texts):
            continue
        boxes = [i for i, t in enumerate(texts) if _is_box(t)]
        if not boxes and len(cells) > 1 and not any(BOX_RE.search(t) for t in texts[1:]) \
                and sum(1 for t in texts[1:] if t and len(t) <= 6) >= 2:
            header_labels = texts          # column titles row: 條件 | 是 | 否 | 理由
            continue
        n += 1
        key = f"{form_id}-{n}"
        if len(boxes) >= 2:
            labels = header_labels if header_labels and len(header_labels) == len(texts) else None
            columns = {(labels[i] if labels else ("是", "否", "N/A")[k] if k < 3 else str(i)): i
                       for k, i in enumerate(boxes)}
            reason = cells[-1] if len(cells) - 1 not in boxes and len(cells) > 2 else None
            subs = BOX_RE.findall(cells[boxes[0]].text)
            question = texts[0]
            if len(subs) > 1:          # one 是/否 pair per sub-statement line
                lines = [p.text.strip() for p in cells[0].paragraphs if p.text.strip()]
                question = lines[0] + "（逐項：" + "／".join(x[:20] for x in lines[-len(subs):]) + "）"
            items.append(Item(key, "yesno", question, cells, tuple(columns), columns, reason))
        elif any(BOX_RE.search(t) for t in texts):
            question = texts[0] if not BOX_RE.match(texts[0]) else ""
            first = cells[0].paragraphs[0].text.strip() if cells[0].paragraphs else ""
            items.append(Item(key, "choice", _clean(first) or question[:80], cells,
                              _options("\n".join(c.text for c in cells))))
        elif len(cells) > 1 and not texts[-1]:
            # label | value pairs (several per row: 編號 | ＿ | 主持人 | ＿)
            pairs = ([(cells[i], cells[i + 1]) for i in range(0, len(cells), 2)]
                     if len(cells) % 2 == 0 and not any(texts[1::2]) else [(cells[0], cells[-1])])
            pairs = [(lab, val) for lab, val in pairs if _clean(lab.text)]
            for k, (lab, val) in enumerate(pairs):
                label = _clean(lab.text).replace(" ", "")
                src = next((v for h, v in HEADER_LABELS.items() if label == h or label.endswith(h)), "")
                question = " ".join(t for t in texts if t) if len(pairs) == 1 else _clean(lab.text)
                items.append(Item(f"{form_id}-{n + k}", "header" if src else "text", question,
                                  [lab, val], answer_cell=val, source=src))
            n += len(pairs) - 1
        else:
            items.append(Item(key, "text", texts[0][:120], cells))
        counter[0] = n
    n = counter[0]
    loose = [p for p in doc.paragraphs if BOX_RE.search(p.text)]
    if loose:                         # checkbox paragraphs outside tables (e.g. FDA checklist)
        items.append(Item(f"{form_id}-{n + 1}", "choice", "勾選項目（表格外）", [],
                          _options("\n".join(p.text for p in loose)), paras=loose))
    return items


# ── answering ───────────────────────────────────────────────────────────────

def _lookup(config, dotted):
    v = config
    for part in dotted.split("."):
        v = v.get(part) if isinstance(v, dict) else None
    return "" if v is None else str(v)


def _lines(value):
    if isinstance(value, list):
        return [str(v).strip() for v in value]
    return [ln.rstrip() for ln in str(value).strip().split("\n")]


def _split(value):
    """'是：理由' → ('是', '理由')."""
    m = re.match(r"\s*([^：:]+?)\s*[：:]\s*(.*)$", str(value), re.S)
    return (m[1], m[2].strip()) if m else (str(value).strip(), "")


def _replace_text(cell, lines):
    """One-cell row: keep the heading paragraph, the answer replaces the guidance below it."""
    paras = [p for p in cell.paragraphs]
    head = next((i for i, p in enumerate(paras) if p.text.strip()), 0)
    body = paras[head + 1:]
    model = body[0] if body else paras[head]
    anchor = paras[head]._p
    for p in body:
        p._p.getparent().remove(p._p)
    for line in lines:
        new = copy.deepcopy(model._p)
        for r in new.findall("{*}r")[1:]:
            new.remove(r)
        anchor.addnext(new)
        anchor = new
        from docx.text.paragraph import Paragraph
        para = Paragraph(new, cell)
        if para.runs:
            para.runs[0].text = line
        else:
            para.add_run(line)


def _write(cell, lines):
    Region(cell.paragraphs).write(lines[0] if lines else "")
    for line in lines[1:]:
        cell.add_paragraph(line)


def _tick_cell(cell, nth=0):
    """Tick the nth box of a cell."""
    for p in cell_paragraphs(cell):
        for r in p.runs:
            for i, ch in enumerate(r.text):
                if ch in UNCHECKED:
                    if nth == 0:
                        r.text = r.text[:i] + check(True) + r.text[i + 1:]
                        return True
                    nth -= 1
    return False


def _tick_option(item, option, note):
    region = Region(item.paragraphs())
    spans, text = region._flat()
    if _box_after(text):           # 是□ 否□ N/A□: the box follows its label
        m = re.search(rf"{re.escape(option)}\s*[{UNCHECKED}]", text)
        if not m:
            return False
        region._replace(spans, m.end() - 1, m.end(), check(True))
    else:
        try:
            region.tick(option)
        except LookupError:
            return False
    if note:
        try:
            region.put(note, after=option)
        except LookupError:
            pass
    return True


def _fill_list(item, value):
    """One answer line per row, cells separated by ｜ (or |): fills empty rows, adds rows as needed."""
    table = item.paras
    rows = [ln for ln in (value if isinstance(value, list) else _lines(value)) if str(ln).strip()]
    width = len(item.options)
    for line in rows:
        cells = [c.strip() for c in re.split(r"[｜|]", str(line))]
        target = None
        for row in table.rows[1:]:
            uc = _unique_cells(row)
            texts = [_clean(c.text) for c in uc]
            if len(uc) != width:
                continue
            # an empty row, or one whose filled cells (e.g. 計畫書 in the first column) match this line
            fits = all(not t or (i < len(cells) and t == cells[i]) for i, t in enumerate(texts))
            room = any(cells[i] and not texts[i] for i in range(min(len(texts), len(cells))))
            if fits and room:
                target = uc
                break
        if target is None:
            new = copy.deepcopy(table.rows[-1]._tr)
            table._tbl.append(new)
            target = _unique_cells(table.rows[-1])
            for c in target:
                Region(c.paragraphs).write("")
        for c, v in zip(target, cells, strict=False):
            if v and not _clean(c.text):
                _write(c, [v])
    return ""


def apply(item, value):
    """Write one answer into its row. Returns a warning string, or ''."""
    if item.kind == "yesno":
        reasons = []
        for nth, entry in enumerate(value if isinstance(value, list) else [value]):
            label, reason = _split(entry)
            col = next((i for name, i in item.columns.items()
                        if name.startswith(label) or label.startswith(name)), None)
            if col is None:
                return f"{item.key}: 「{label}」不是 {'/'.join(item.columns)}"
            _tick_cell(item.cells[col], nth)
            if reason:
                reasons.append(reason)
        if reasons and item.answer_cell is not None:
            _write(item.answer_cell, reasons)
        return ""
    if item.kind == "list":
        return _fill_list(item, value)
    if item.kind == "choice":
        bad = []
        for entry in (value if isinstance(value, list) else [value]):
            option, note = _split(entry)
            match = next((o for o in item.options if o.startswith(option) or option.startswith(o)), None)
            if not match or not _tick_option(item, match[:20], note):
                bad.append(option)
        return f"{item.key}: 找不到選項 {'、'.join(bad)}" if bad else ""
    lines = _lines(value)
    if item.answer_cell is not None:
        _write(item.answer_cell, lines)
    else:
        _replace_text(item.cells[0], lines)
    return ""


def fill_inline_headers(items, config):
    """'中文：' / '試驗主持人：' inside a one-cell header row (consent forms)."""
    for item in items[:4]:
        if item.kind != "text" or item.answer_cell is not None:
            continue
        region = Region([p for c in item.cells for p in cell_paragraphs(c)])
        for label, src in INLINE_LABELS:
            value = src(config) if callable(src) else _lookup(config, src)
            if value and label in region._flat()[1]:
                try:
                    region.put(value, after=label)
                except LookupError:
                    pass


def _pi_unit(config):
    return str(config.get("pi", {}).get("dept", "")).partition("／")[0].strip()


PARA_LABELS = (            # body paragraphs outside tables: 'label' followed by blanks → value
    ("倫委會案件編號：", "study.irb_no"), ("倫委會案號（REC Number）：", "study.irb_no"),
    ("計畫名稱：", "study.title_zh"), ("計畫名稱（Study Title）：", "study.title_zh"),
    ("計畫主持人（Principal investigator）：", "pi.name"), ("本計畫主持人", "pi.name"),
    ("立切結書人姓名正楷：", "pi.name"), ("服務單位：", _pi_unit), ("職稱：", _pi_title),
    ("執行", "study.title_zh"),       # 本計畫主持人＿＿執行＿＿研究計畫
)
# the label's blank is empty: only spaces/underscores until the line ends or the next 'label：'
BLANK_AFTER = re.compile(r"[\s＿_　]*(?:$|[^\s＿_　：:]{1,8}[：:])|[\s　]*[＿_]{2,}")


def fill_paragraph_labels(doc, config):
    """Fill 'label：＿＿' lines in the body (not tables) whose blank is still empty."""
    for para in doc.paragraphs:
        text = para.text
        for label, src in PARA_LABELS:
            i = text.find(label)
            if i < 0:
                continue
            value = src(config) if callable(src) else _lookup(config, src)
            if not value or not BLANK_AFTER.match(text, i + len(label)):
                continue
            try:
                Region([para]).put(value, after=label)
            except LookupError:
                pass
            text = para.text


def fill(doc, form_id, config, answers, defaults=None):
    """Fill a scanned blank in place. answers: {key: value}; defaults: {question prefix: callable(config)}."""
    items = scan(doc, form_id)
    warnings = []
    fill_inline_headers(items, config)
    fill_paragraph_labels(doc, config)
    for item in items:
        value = answers.get(item.key)
        if value in (None, "", []):
            value = next((fn(config) for prefix, fn in (defaults or {}).items()
                          if item.question.replace(" ", "").startswith(prefix)), None)
        if value in (None, "", []) and item.kind == "header":
            value = _lookup(config, item.source)
        if value in (None, "", []):
            continue
        warn = apply(item, value)
        if warn:
            warnings.append(warn)
    return items, warnings


def skeleton_md(items, title):
    """Markdown answer template for one blank."""
    out = [f"<!-- ════ {title} ════ -->", ""]
    for it in items:
        if it.kind == "header":
            continue
        out.append(f"## {it.key} {it.question[:70]}")
        if it.kind == "yesno":
            multi = "（逐項" in it.question
            out.append(f"<!-- {' / '.join(it.columns)}" + ("；逐項各一行「- 是」" if multi else "")
                       + ("；可寫「是：理由」" if it.answer_cell is not None else "") + " -->")
        elif it.kind == "list":
            out.append(f"<!-- 清單：每列一行，欄位以「｜」分隔：{' ｜ '.join(it.options)} -->")
        elif it.kind == "choice":
            out.append("<!-- 勾選：每項一行「- 選項開頭」，可複選；選項：" + " / ".join(o[:25] for o in it.options) + " -->")
        else:
            out.append("<!-- 文字；不填就保留範本原文 -->")
        out.append("")
    return "\n".join(out)


# ── rech attachments --skeleton ─────────────────────────────────────────────

ANSWERS_FILE = "附件填寫.md"


def row_forms(config, extra=()):
    """[(form id, name)] of the study's forms whose generator module declares them in ROW_FORMS."""
    import importlib

    from rec_helper.form_selector import _pack, select_forms
    registry = _pack().FORM_REGISTRY
    chosen = [fid for fid, _ in select_forms(config)] + [f for f in extra if f in registry]
    out = []
    for fid in dict.fromkeys(chosen):
        module = registry.get(fid, (None, ""))[1]
        if module.startswith("generators."):       # shared rebuild generators: never row-filled
            continue
        if module and fid in getattr(importlib.import_module(module), "ROW_FORMS", {}):
            out.append((fid, registry[fid][0]))
    return out


def main(config_path="config.toml", phase=None, skeleton=False, forms=(), path=ANSWERS_FILE):
    """List (or append to 附件填寫.md) the questions of the study's row-filled attachments."""
    import os

    from docx import Document

    from rec_helper.config import ConfigError, load_config
    from rec_helper.institution import activate, current
    try:
        config = load_config(config_path, phase=phase)
    except ConfigError as e:
        print(f"✗ {e}")
        return 1
    activate(config)
    given = {k.split()[0] for k in (config.get("attachments") or {}) if k.strip()} \
        if isinstance(config.get("attachments"), dict) else set()
    chunks = []
    found = row_forms(config, forms)
    if not found:
        print("  本案這個階段沒有逐題填寫的附件（可用 --form <ID> 指定，例如 --form DSMP）")
    for fid, name in found:
        blank = current().blank_path(fid)
        if not os.path.isfile(blank):
            print(f"  ✗ {fid} {name}: 找不到空白表單 {blank}")
            continue
        items = [it for it in scan(Document(blank), fid) if it.key not in given]
        todo = [it for it in items if it.kind != "header"]
        print(f"  ■ {fid} {name}：{len(todo)} 題未作答")
        if skeleton and todo:
            chunks.append(skeleton_md(items, f"{name}（{fid}）"))
    if skeleton and chunks:
        new = not os.path.exists(path)
        with open(path, "a", encoding="utf-8") as fh:
            if new:
                fh.write("# 附件填寫\n\n<!-- config.toml: attachments = \"@附件填寫.md\"；"
                         "每題 `## <題號>`，不填的題目保留範本原文 -->\n\n")
            fh.write("\n\n" + "\n\n".join(chunks) + "\n")
        print(f"■ {path} {'written' if new else 'appended'} — 填好後 `rech generate`"
              + ("；記得在 config.toml 加 attachments = \"@附件填寫.md\"" if new else ""))
    return 0
