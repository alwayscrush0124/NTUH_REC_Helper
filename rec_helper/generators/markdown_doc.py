"""Render a study's own Markdown (計畫書.md, 英文摘要.md) as a submission-ready DOCX.

Supported: `# 標題` (centered title), `## 章` / `### 節` headings, paragraphs
(wrapped lines joined, no space next to CJK), `-`/`*` bullets, `1.` numbered
items, `**bold**` (rendered plain-bold), `<!-- comments -->` (dropped), and
`【待寫】` placeholders, which are kept and highlighted so nothing unfinished
is submitted unnoticed.
"""
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt

from rec_helper.docx_utils import init_doc, set_run_font

TODO_MARK = "【待寫】"
CJK = r"[　-〿一-鿿＀-￯]"


def _join(lines):
    out = ""
    for ln in (x.strip() for x in lines):
        if out and not re.search(CJK + "$", out) and not re.match(CJK, ln):
            out += " "
        out += ln
    return out


def _runs(paragraph, text, size, bold=False):
    """Add text with **bold** spans and highlighted 【待寫】."""
    for part in re.split(r"(\*\*[^*]+\*\*|【待寫】[^\n]*)", text):
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if is_bold else part)
        set_run_font(run, None, size, bold or is_bold)
        if part.startswith(TODO_MARK):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def blocks(markdown):
    """[(kind, text)] with kind in title, h2, h3, bullet, number, para."""
    text = re.sub(r"<!--.*?-->", "", markdown, flags=re.S)
    out, para = [], []

    def flush():
        if para:
            out.append(("para", _join(para)))
            para.clear()

    for line in text.splitlines():
        s = line.strip()
        if not s:
            flush()
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            flush()
            out.append(({1: "title", 2: "h2", 3: "h3"}[len(m.group(1))], m.group(2).strip()))
        elif re.match(r"^[-*]\s+", s):
            flush()
            out.append(("bullet", re.sub(r"^[-*]\s+", "", s)))
        elif re.match(r"^\d+[.、)]\s*", s):
            flush()
            out.append(("number", s))
        elif out and out[-1][0] in ("bullet", "number") and line.startswith(("  ", "\t")):
            out[-1] = (out[-1][0], _join([out[-1][1], s]))
        else:
            para.append(s)
    flush()
    return out


def render(markdown, path, header=None, size=12):
    """Write the Markdown to `path` as DOCX (institution font); `header` goes top-right."""
    doc = init_doc(size)
    if header:
        from rec_helper.versions import add_header
        add_header(doc, header)
    for kind, text in blocks(markdown):
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        if kind == "title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(p, text, size + 4, bold=True)
            fmt.space_after = Pt(12)
        elif kind in ("h2", "h3"):
            _runs(p, text, size + (1 if kind == "h2" else 0), bold=True)
            fmt.space_before, fmt.space_after = Pt(10 if kind == "h2" else 6), Pt(4)
            fmt.keep_with_next = True
        elif kind == "bullet":
            _runs(p, "• " + text, size)
            fmt.left_indent, fmt.first_line_indent = Pt(18), Pt(-10)
        elif kind == "number":
            _runs(p, text, size)
            fmt.left_indent, fmt.first_line_indent = Pt(18), Pt(-14)
        else:
            _runs(p, text, size)
            fmt.first_line_indent = Pt(size * 2)
            fmt.space_after = Pt(4)
    doc.save(path)
    return path


def todo_count(markdown):
    return len(re.findall(TODO_MARK, re.sub(r"<!--.*?-->", "", markdown, flags=re.S)))
