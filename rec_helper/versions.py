"""Version numbers and dates on submitted documents, and a ledger of what was sent.

The committee wants a 版本日期 on the protocol, summaries, consent forms,
questionnaires and recruitment material, and an amendment cites the approved
version it changes. One place holds them, config.toml:

    [version]
    number = "1.0"
    date   = "2026/10/14"          # default: today
    計畫書  = "1.1"                  # per document: a number, or { number = "1.1", date = "…" }

Generated documents (計畫書, 中文/英文摘要) print it in their header. Documents
written elsewhere get it with `rech stamp FILE…` — a stamped *copy* in
output/stamped/ (the original is never touched) — and every stamp is appended
to 版本紀錄.md with the file's SHA-256, so an approved version can be identified later.
"""
import datetime
import hashlib
import os
import re

LEDGER = "版本紀錄.md"
STAMP_DIR = "stamped"


def _norm_date(value):
    m = re.match(r"\s*(\d{4})\D+(\d{1,2})\D+(\d{1,2})", str(value or ""))
    return f"{m[1]}/{int(m[2]):02d}/{int(m[3]):02d}" if m else datetime.date.today().strftime("%Y/%m/%d")


def version_of(config, doc=None):
    """(number, date) for a document name, falling back to [version] number/date."""
    v = config.get("version") or {}
    number, date = str(v.get("number", "1.0")), v.get("date")
    own = v.get(doc) if doc else None
    if isinstance(own, dict):
        number, date = str(own.get("number", number)), own.get("date", date)
    elif own:
        number = str(own)
    return number, _norm_date(date)


def label(config, doc=None, english=False):
    number, date = version_of(config, doc)
    return f"Version {number}, Date: {date}" if english else f"版本：{number}　版本日期：{date}"


def add_header(doc, text):
    """`text` right-aligned, 9 pt, in every section header of a python-docx Document."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        p = header.add_paragraph() if header.paragraphs and header.paragraphs[0].text.strip() else \
            (header.paragraphs[0] if header.paragraphs else header.add_paragraph())
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(text)
        run.font.size = Pt(9)
    return doc


def stamp_docx(path, out, text):
    """Copy a DOCX with `text` right-aligned in every section header."""
    from docx import Document
    add_header(Document(path), text).save(out)
    return out


def _cjk_font():
    for f in (r"C:\Windows\Fonts\kaiu.ttf", "/Library/Fonts/Kaiti.ttc", "/System/Library/Fonts/Supplemental/Songti.ttc",
              "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"):
        if os.path.isfile(f):
            return f
    return None


def stamp_pdf(path, out, text):
    """Copy a PDF with `text` in the top-right margin of every page (needs PyMuPDF)."""
    try:
        import pymupdf
    except ImportError as e:
        raise RuntimeError("stamping PDFs needs PyMuPDF: pip install pymupdf (or stamp the Word file instead)") from e
    font = _cjk_font()
    with pymupdf.open(path) as doc:
        for page in doc:
            r = page.rect
            box = pymupdf.Rect(r.x0 + 40, 18, r.x1 - 45, 36)
            kw = {"fontname": "kai", "fontfile": font} if font else {"fontname": "china-t"}
            page.insert_textbox(box, text, fontsize=9, align=2, **kw)
        if font:
            doc.subset_fonts()
        doc.save(out, garbage=4, deflate=True)
    return out


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 16), b""):
            h.update(block)
    return h.hexdigest()


def record(rows, ledger=LEDGER):
    """Append [(file, number, date, phase)] to the ledger (created with a header)."""
    new = not os.path.exists(ledger)
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    with open(ledger, "a", encoding="utf-8") as fh:
        if new:
            fh.write("# 版本紀錄\n\n<!-- rech stamp / rech versions --record 自動附加；SHA-256 可比對送審或核准的檔案 -->\n\n"
                     "| 記錄時間 | 階段 | 檔案 | 版本 | 版本日期 | SHA-256 |\n|---|---|---|---|---|---|\n")
        for path, number, date, phase in rows:
            fh.write(f"| {now} | {phase} | {os.path.basename(path)} | {number} | {date} | `{sha256(path)[:16]}` |\n")


def stamp(files, config, output_dir="output", number=None, date=None, doc=None):
    """Stamp each file into output/stamped/; returns the stamped paths."""
    out_dir = os.path.join(output_dir, STAMP_DIR)
    os.makedirs(out_dir, exist_ok=True)
    n0, d0 = version_of(config, doc)
    number, date = number or n0, _norm_date(date or d0)
    done = []
    for path in files:
        ext = os.path.splitext(path)[1].lower()
        english = bool(re.search(r"英文|english|abstract", os.path.basename(path), re.I))
        text = f"Version {number}, Date: {date}" if english else f"版本：{number}　版本日期：{date}"
        out = os.path.join(out_dir, os.path.basename(path))
        if ext == ".pdf":
            stamp_pdf(path, out, text)
        elif ext == ".docx":
            stamp_docx(path, out, text)
        else:
            print(f"  ✗ {path}: only .pdf / .docx can be stamped")
            continue
        done.append(out)
        print(f"  ■ {out}  ({text})")
    if done:
        record([(p, number, date, config.get("phase", "")) for p in done])
    return done


def main(files, config_path="config.toml", output_dir="output", number=None, date=None, doc=None, record_outputs=False):
    from rec_helper.config import ConfigError, load_config
    try:
        config = load_config(config_path)
    except ConfigError as e:
        print(f"✗ {e}")
        return 2
    if record_outputs:
        pdfs = sorted(os.path.join(output_dir, f) for f in os.listdir(output_dir) if f.endswith(".pdf")) \
            if os.path.isdir(output_dir) else []
        n, d = version_of(config, doc)
        record([(p, number or n, _norm_date(date or d), config.get("phase", "")) for p in pdfs])
        print(f"■ {len(pdfs)} PDF 記入 {LEDGER}")
        return 0
    missing = [f for f in files if not os.path.isfile(f)]
    if missing:
        print(f"✗ not found: {', '.join(missing)}")
        return 1
    try:
        stamp(files, config, output_dir, number, date, doc)
    except RuntimeError as e:
        print(f"✗ {e}")
        return 1
    print(f"■ 原始檔未更動；版本紀錄 → {LEDGER}")
    return 0
