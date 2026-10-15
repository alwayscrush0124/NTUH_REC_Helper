"""Everything a study says, gathered in one place for the cross-checks.

consistency.py, precheck.py, guidelines.py and refs.py all read the study the
same way: the structured config, the proposal summary, the web-form answers
(given in 線上填寫.md *or* derived from config), the attachment answers, and
the protocol — either `protocol = "計畫書.md"` (rendered by `rech protocol`) or
an existing document named by `protocol_file = "…docx|pdf|md"`.

Paths are relative to the study folder (the cwd), like every other study file.
"""
import copy
import os
import re

PROTOCOL_MD = "計畫書.md"


def read_document(path):
    """Plain text of a .md/.txt/.docx/.pdf file ('' when it can't be read)."""
    if not path or not os.path.isfile(path):
        return ""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".md", ".txt"):
        with open(path, encoding="utf-8") as fh:
            return re.sub(r"<!--.*?-->", "", fh.read(), flags=re.S)
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
        return "\n".join(parts)
    if ext == ".pdf":
        return _pdf_text(path)
    return ""


def _pdf_text(path):
    try:
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    except ImportError:
        pass
    try:
        import pymupdf
        with pymupdf.open(path) as doc:
            return "\n".join(page.get_text() for page in doc)
    except ImportError:
        return ""


def protocol_path(config):
    """The protocol to check: protocol_file, else the 計畫書.md source, else None."""
    for key in ("protocol_file", "protocol"):
        value = config.get(key)
        if isinstance(value, str) and value and os.path.isfile(value):
            return value
    return PROTOCOL_MD if os.path.isfile(PROTOCOL_MD) else None


def protocol_text(config):
    return read_document(protocol_path(config))


def _flat(value):
    if isinstance(value, dict):
        return "\n".join(f"{k}\n{_flat(v)}" for k, v in value.items())
    if isinstance(value, list):
        return "\n".join(_flat(v) for v in value)
    return "" if value is None else str(value)


def proposal_text(config):
    """The summary with its Chinese headings back (sections are keyed background, design, …)."""
    from rec_helper.config import SECTION_KEYS
    p = config.get("proposal")
    if isinstance(p, dict):
        heading = {v: k for k, v in SECTION_KEYS.items()}
        return _flat({heading.get(k, k): v for k, v in p.items()})
    return _flat(p) if p else ""


def online_answers(config, phase="new"):
    """{field id: answer} for the web forms of `phase`, from 線上填寫.md or derived from config."""
    from rec_helper import online_forms
    c = copy.deepcopy(config)
    c["phase"] = phase
    given = online_forms.overrides(c)
    out = {}
    for form in online_forms.online_forms(c):
        filled, _ = online_forms.fill(form, c, given)
        for _, rows in filled:
            for fld, ans, src in rows:
                if src != "todo" and ans not in ("", []):
                    out[fld.id] = ans
    return out


def attachment_answers(config):
    raw = config.get("attachments")
    return {k.split()[0]: v for k, v in raw.items() if k.strip()} if isinstance(raw, dict) else {}


def corpus(config):
    """{source name: text} — what reviewers read, per document."""
    online = online_answers(config)
    return {
        "中文計畫摘要": proposal_text(config),
        "線上申請書": "\n".join(f"{k}: {_flat(v)}" for k, v in online.items()),
        "附件填寫": "\n".join(f"{k}: {_flat(v)}" for k, v in attachment_answers(config).items()),
        "計畫書": protocol_text(config),
    }
