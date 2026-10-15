"""Question-per-row blanks (rec_helper.form_rows) and the NTUH routing built on them."""
import os

import pytest
from docx import Document

from rec_helper import form_rows

CONFIG = {"study": {"title_zh": "示範研究", "title_en": "Demo", "irb_no": ""},
          "pi": {"name": "王大明", "dept": "胸腔內科部／醫師"}, "co_pi": [{"name": "李小華"}]}


def _blank(tmp_path):
    doc = Document()
    doc.add_paragraph("計畫名稱：")
    t = doc.add_table(rows=6, cols=4)
    rows = [("試驗主持人", "", "", ""), ("必須符合以下所有條件：", "是", "否", "請說明理由"),
            ("已進行臨床前試驗？", "□", "□", ""), ("風險等級：□1.最小風險 □2.微幅超過", "", "", ""),
            ("請說明監測方式", "", "", ""), ("收藥紀錄", "是□　否□　N/A□", "", "")]
    for r, values in zip(t.rows, rows, strict=False):
        for c, v in zip(r.cells, values, strict=False):
            c.text = v
    for r in (3, 4):                        # one-cell rows: merge the question across the row
        t.rows[r].cells[0].merge(t.rows[r].cells[3])
    t.rows[5].cells[1].merge(t.rows[5].cells[3])
    log = doc.add_table(rows=4, cols=3)
    for c, v in zip(log.rows[0].cells, ("序號", "日期", "事件"), strict=False):
        c.text = v
    path = tmp_path / "blank.docx"
    doc.save(path)
    return path


def test_scan_classifies_rows(tmp_path):
    items = form_rows.scan(Document(_blank(tmp_path)), "X")
    kinds = {it.question[:6]: it.kind for it in items}
    assert kinds["試驗主持人"] == "header"
    assert kinds["已進行臨床前"] == "yesno"
    assert kinds["風險等級：□"] == "choice"
    assert kinds["序號／日期／"] == "list"
    drug = next(it for it in items if it.question.startswith("收藥紀錄"))
    assert drug.options == ("是", "否", "N/A")          # 'label□' order kept


def test_fill_ticks_writes_and_logs(tmp_path):
    doc = Document(_blank(tmp_path))
    keys = {it.question[:4]: it.key for it in form_rows.scan(doc, "X")}
    answers = {keys["已進行臨"]: "是：已有動物試驗", keys["風險等級"]: "1.最小風險",
               keys["收藥紀錄"]: "N/A", keys["請說明監"]: "每三個月內部監測",
               keys["序號／日"]: ["1｜2026/01/02｜發燒住院", "2｜2026/03/04｜跌倒"]}
    _, warnings = form_rows.fill(doc, "X", CONFIG, answers)
    assert not warnings
    text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "王大明" in text and "■" in text and "已有動物試驗" in text
    assert "■1.最小風險" in text and "N/A■" in text and "否□" in text
    assert "每三個月內部監測" in text and "跌倒" in text
    assert "示範研究" in doc.paragraphs[0].text               # body 'label：' line


def test_skeleton_lists_questions(tmp_path):
    md = form_rows.skeleton_md(form_rows.scan(Document(_blank(tmp_path)), "X"), "示範表")
    assert "## X-2 已進行臨床前試驗？" in md and "## X-1 " not in md and "｜" in md
    assert "試驗主持人" not in md                         # headers come from config


@pytest.fixture
def ntuh(monkeypatch):
    monkeypatch.setenv("REC_INSTITUTION", "ntuh")


def test_ntuh_output_names_map_back_and_are_uploadable(ntuh):
    from rec_helper.institution import current
    from rec_helper.institutions.ntuh import attachments as at
    for fid in at.ROW_FORMS:
        name = os.path.basename(at.row_form_path(fid, "."))
        assert current().form_id(name) == fid, name
        assert not any(ch in name for ch in "()\"'<>|") and "--" not in name   # e-REC upload rules


def test_ntuh_every_phase_routes_to_registered_forms(ntuh):
    from rec_helper.config import load_config
    from rec_helper.form_selector import load_generator, select_forms
    from rec_helper.institutions.ntuh.forms import PHASE_FORMS
    config = load_config("config.toml")
    config["institution"] = "ntuh"
    for phase in PHASE_FORMS:
        config["phase"] = phase
        for fid, _ in select_forms(config):
            assert callable(load_generator(fid)), (phase, fid)


def test_ntuh_branch_consent_follows_the_pi(ntuh):
    from rec_helper.institutions.ntuh.forms import site
    assert site({"pi": {"dept": "新竹臺大分院影像醫學部／醫事放射師"}}) == "HCH"
    assert site({"pi": {"dept": "內科部／醫師"}}) == ""
    assert site({"site": "bmh", "pi": {}}) == "BMH"
