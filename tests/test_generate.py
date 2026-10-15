"""Generation end to end: config → selected forms → DOCX (NTUH pack)."""
import os

import pytest
from docx import Document

from rec_helper import generate_all
from rec_helper.config import load_config
from rec_helper.form_selector import FORM_REGISTRY, get_generator

EXAMPLE = "rec_helper/examples/ntuh-chart-review/config.toml"
BLANKS = os.path.abspath("templates")
HAVE_BLANKS = os.path.isdir(os.path.join(BLANKS, "ntuh"))


@pytest.fixture
def config(monkeypatch):
    monkeypatch.setenv("REC_INSTITUTION", "ntuh")
    monkeypatch.setenv("REC_TEMPLATES", BLANKS)
    return load_config(EXAMPLE)


@pytest.fixture
def output_dir(tmp_path):
    path = tmp_path / "output"
    path.mkdir()
    return str(path)


def docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


@pytest.mark.skipif(not HAVE_BLANKS, reason="NTUH blanks not downloaded (rech templates)")
@pytest.mark.parametrize("form_id", sorted(FORM_REGISTRY))
def test_every_registered_generator_runs(form_id, config, output_dir):
    """Each registry entry points at an importable function that writes a DOCX."""
    assert get_generator(form_id) is not None
    path = generate_all.generate_form(form_id, config, output_dir)
    assert os.path.getsize(path) > 0


@pytest.mark.skipif(not HAVE_BLANKS, reason="NTUH blanks not downloaded (rech templates)")
def test_attachments_carry_the_study(config, output_dir):
    assert "王小明" in docx_text(generate_all.generate_form("MEDREC", config, output_dir))
    assert "SGLT2" in docx_text(generate_all.generate_form("AF164", config, output_dir))


def test_generate_all_reports_bad_config(tmp_path, capsys):
    bad = tmp_path / "config.toml"
    bad.write_text('phase = "new"\n[study]\n[pi]\n', encoding="utf-8")
    assert generate_all.main(str(bad), str(tmp_path / "out")) == 2
    assert "study.title_zh" in capsys.readouterr().out
    assert generate_all.main(str(tmp_path / "missing.toml"), str(tmp_path / "out")) == 2


def test_checklist_generation(config, tmp_path):
    from rec_helper.checklist import generate_checklist
    results = [
        ("AF001", "臨床研究計畫申請人自評表", "/fake/path.docx", "generated"),
        ("MEDREC", "使用病歷資料進行研究申請表", None, "error"),
    ]
    path = str(tmp_path / "checklist.md")
    generate_checklist(config, results, "新案審查", path)
    with open(path, encoding="utf-8") as f:
        content = f.read()
    assert "■ AF001" in content and "□ MEDREC" in content
    assert "SGLT2" in content


def test_proposal_summary_uses_config_text(config, output_dir):
    """proposal.* text fills 中文計畫摘要; absent keys keep the placeholder."""
    from rec_helper.generators.proposal import generate_proposal_summary
    text = "\n".join(p.text for p in Document(generate_proposal_summary(config, output_dir)).paragraphs)
    assert "急性腎損傷" in text and "1. 比較 SGLT2" in text
    assert "（請列出）" not in text
    config.pop("proposal")
    text = "\n".join(p.text for p in Document(generate_proposal_summary(config, output_dir)).paragraphs)
    assert "納入條件：（請列出）" in text
